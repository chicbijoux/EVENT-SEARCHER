import sys, os, json, queue, threading, time, csv, io
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

try:
    from dotenv import load_dotenv
    load_dotenv(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".env"), override=True)
except ImportError:
    pass

from flask import Flask, Response, jsonify, render_template, request, stream_with_context
from email_scraper import Scraper, init_db, save_events, CSV_PATH, enrich_summaries

app = Flask(__name__)


class WebScraper(Scraper):
    def __init__(self, emit=None, **kwargs):
        # Disable Playwright in threads — it requires its own event loop
        kwargs.setdefault('use_playwright', False)
        super().__init__(**kwargs)
        self._emit = emit or (lambda _: None)

    def _send(self, **kwargs):
        self._emit(kwargs)

    def ddg_search(self, query: str, num_results: int = 10) -> list[str]:
        try:
            from ddgs import DDGS
        except ImportError:
            self._send(type='log', msg='ddgs not installed — run: pip install ddgs', cls='err')
            return []
        urls = []
        try:
            with DDGS() as ddgs:
                for r in ddgs.text(query, max_results=num_results):
                    u = r.get('href', '')
                    if u:
                        urls.append(u)
        except Exception as e:
            self._send(type='log', msg=f'Search error: {e}', cls='warn')
        return urls

    def search_events_streaming(self, cities, queries_per_city=2, results_per_query=5, custom_queries=None, topic=''):
        all_events = []
        seen = set()
        total_urls = []

        if custom_queries:
            default_city = cities[0] if cities else 'Cluj-Napoca'

            # Use the exact user input as the relevance filter (split on spaces for multi-word topics)
            import re as _re
            topic_kws = [w for w in _re.split(r'\s+', topic.lower()) if len(w) > 1] if topic else []

            self._send(type='log', msg=f'▶ Running {len(custom_queries)} custom queries', cls='step')
            if topic_kws:
                self._send(type='log', msg=f'✦ Topic filter: {", ".join(topic_kws)}', cls='info')
            self._send(type='step', n=1, state='active', badge='searching…')

            url_set = set()
            for i, q in enumerate(custom_queries):
                self._send(type='log', msg=f'› {q}', cls='info')
                for u in self.ddg_search(q, results_per_query):
                    if u not in url_set:
                        url_set.add(u)
                        total_urls.append(u)
                time.sleep(1.0 if i % 5 == 4 else 0.5)

            self._send(type='step', n=1, state='done', badge=f'{len(total_urls)} URLs')
            self._send(type='stats', urls=len(total_urls), crawled=0, events=0, emails=0)
            self._send(type='step', n=2, state='active', badge='crawling…')

            crawled = 0
            for url in total_urls:
                html = self.fetch(url)
                if not html:
                    continue
                crawled += 1
                for ev in self.extract_events(html, url, default_city, family_filter=False):
                    key = ev['event_name'].lower().strip()
                    if key in seen or len(key) <= 3:
                        continue
                    # Keep only events relevant to the searched topic
                    ev_text = (ev.get('event_name', '') + ' ' + ev.get('summary', '')).lower()
                    if topic_kws and not any(kw in ev_text for kw in topic_kws):
                        continue
                    seen.add(key)
                    all_events.append(ev)
                    self._send(type='event', event=ev)
                    self._send(type='log',
                               msg=f'+ {ev["event_name"]} ({ev.get("date", "")})', cls='ok')
                pct = 15 + int(crawled / max(len(total_urls), 1) * 65)
                self._send(type='progress', pct=pct)
                self._send(type='stats', urls=len(total_urls), crawled=crawled,
                           events=len(all_events),
                           emails=sum(1 for e in all_events if e.get('email')))
                time.sleep(self.delay)

            self._send(type='step', n=2, state='done', badge=f'{crawled} crawled')

        else:
            # Original city-based template flow
            for city in cities:
                self._send(type='log', msg=f'▶ Searching: {city}', cls='step')
                self._send(type='step', n=1, state='active', badge='searching…')

                city_urls = []
                for tmpl in self.EVENT_QUERIES[:queries_per_city]:
                    q = tmpl.format(city=city)
                    self._send(type='log', msg=f'› {q}', cls='info')
                    for u in self.ddg_search(q, results_per_query):
                        if u not in city_urls:
                            city_urls.append(u)
                    time.sleep(0.5)

                total_urls.extend(city_urls)
                self._send(type='step', n=1, state='done', badge=f'{len(total_urls)} URLs')
                self._send(type='stats', urls=len(total_urls), crawled=0,
                           events=len(all_events), emails=sum(1 for e in all_events if e.get('email')))
                self._send(type='log', msg=f'✓ {len(city_urls)} URLs for {city}', cls='ok')

                self._send(type='step', n=2, state='active', badge='crawling…')
                crawled = 0
                for url in city_urls:
                    html = self.fetch(url)
                    if not html:
                        continue
                    crawled += 1
                    for ev in self.extract_events(html, url, city):
                        key = ev['event_name'].lower().strip()
                        if key not in seen and len(key) > 3:
                            seen.add(key)
                            all_events.append(ev)
                            self._send(type='event', event=ev)
                            self._send(type='log',
                                       msg=f'+ {ev["event_name"]} ({ev.get("date", "")})', cls='ok')
                    pct = 15 + int(crawled / max(len(city_urls), 1) * 65)
                    self._send(type='progress', pct=pct)
                    self._send(type='stats', urls=len(total_urls), crawled=crawled,
                               events=len(all_events),
                               emails=sum(1 for e in all_events if e.get('email')))
                    time.sleep(self.delay)

                self._send(type='step', n=2, state='done', badge=f'{crawled} crawled')

        return all_events


@app.route('/')
def index():
    return render_template('index.html', cities=Scraper.CITIES)


@app.route('/api/search')
def search_stream():
    cities = [c.strip() for c in request.args.get('cities', 'Cluj-Napoca').split(',') if c.strip()]
    queries = min(int(request.args.get('queries', 4)), 4)
    results = min(int(request.args.get('results', 50)), 50)
    cq_raw = request.args.get('cq', '')
    custom_queries = json.loads(cq_raw) if cq_raw else None
    topic = request.args.get('topic', '').strip().lower()

    q = queue.Queue()

    def run():
        try:
            scraper = WebScraper(emit=q.put, timeout=15, delay=1.0)
            q.put({'type': 'log',
                   'msg': f'Pipeline started — {len(cities)} cities × {queries} queries × {results} results/query',
                   'cls': 'step'})
            q.put({'type': 'progress', 'pct': 5})

            events = scraper.search_events_streaming(cities, queries, results, custom_queries=custom_queries, topic=topic)

            q.put({'type': 'step', 'n': 3, 'state': 'active', 'badge': 'saving…'})
            if events:
                events = enrich_summaries(events, emit=q.put)
                con = init_db()
                try:
                    save_events(con, events)
                    q.put({'type': 'log',
                           'msg': f'✅ {len(events)} events saved to DB + CSV', 'cls': 'ok'})
                except PermissionError:
                    # DB save still worked; only CSV was blocked
                    from email_scraper import _append_events_csv
                    q.put({'type': 'log',
                           'msg': '⚠ CSV locked (close it in Excel). Events saved to DB only.', 'cls': 'warn'})
                finally:
                    con.close()
                q.put({'type': 'step', 'n': 3, 'state': 'done', 'badge': f'{len(events)} saved'})
            else:
                q.put({'type': 'step', 'n': 3, 'state': 'done', 'badge': '0 saved'})
                q.put({'type': 'log',
                       'msg': 'No events found. Try more cities or queries.', 'cls': 'warn'})

            q.put({'type': 'progress', 'pct': 100})
            q.put({'type': 'done', 'total': len(events)})
        except Exception as e:
            q.put({'type': 'log', 'msg': f'Error: {e}', 'cls': 'err'})
            q.put({'type': 'done', 'total': 0})
        finally:
            q.put(None)

    threading.Thread(target=run, daemon=True).start()

    def generate():
        while True:
            try:
                item = q.get(timeout=180)
                if item is None:
                    yield 'data: {"type":"end"}\n\n'
                    break
                yield f'data: {json.dumps(item, ensure_ascii=False)}\n\n'
            except queue.Empty:
                yield 'data: {"type":"ping"}\n\n'

    return Response(
        stream_with_context(generate()),
        mimetype='text/event-stream',
        headers={'Cache-Control': 'no-cache', 'X-Accel-Buffering': 'no'}
    )


@app.route('/api/events')
def get_events():
    con = init_db()
    rows = con.execute(
        "SELECT event_name, city, date, webpage, email, phone, summary, timestamp "
        "FROM events ORDER BY timestamp DESC"
    ).fetchall()
    con.close()
    keys = ['event_name', 'city', 'date', 'webpage', 'email', 'phone', 'summary', 'timestamp']
    return jsonify([dict(zip(keys, r)) for r in rows])


@app.route('/api/events', methods=['DELETE'])
def clear_events():
    con = init_db()
    con.execute("DELETE FROM events")
    con.commit()
    con.close()
    if os.path.exists(CSV_PATH):
        os.remove(CSV_PATH)
    return jsonify({'ok': True})


@app.route('/api/export/csv')
def export_csv():
    con = init_db()
    rows = con.execute(
        "SELECT event_name, city, date, webpage, email, phone, summary, timestamp "
        "FROM events ORDER BY city, date"
    ).fetchall()
    con.close()
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow(['Eveniment', 'Oraș', 'Dată', 'Website', 'Email', 'Telefon', 'Descriere', 'Salvat la'])
    w.writerows(rows)
    out.seek(0)
    return Response(
        '\ufeff' + out.getvalue(),
        mimetype='text/csv; charset=utf-8',
        headers={'Content-Disposition': 'attachment; filename=events_export.csv'}
    )


@app.route('/api/export/xlsx')
def export_xlsx():
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    con = init_db()
    rows = con.execute(
        "SELECT event_name, city, date, webpage, email, phone, summary, timestamp "
        "FROM events ORDER BY city, date"
    ).fetchall()
    con.close()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Evenimente"

    headers = ['Eveniment', 'Oraș', 'Dată', 'Website', 'Email', 'Telefon', 'Descriere', 'Salvat la']
    col_widths = [45, 16, 14, 40, 32, 18, 60, 18]

    header_font = Font(name='Calibri', bold=True, color='FFFFFF', size=11)
    header_fill = PatternFill(fill_type='solid', fgColor='1F3864')
    header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    thin = Side(style='thin', color='CCCCCC')
    cell_border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # Write headers
    for col, (h, w) in enumerate(zip(headers, col_widths), 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = cell_border
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.row_dimensions[1].height = 22

    # Alternating row fills
    fill_even = PatternFill(fill_type='solid', fgColor='EEF2F7')
    fill_odd  = PatternFill(fill_type='solid', fgColor='FFFFFF')
    link_font = Font(name='Calibri', color='1155CC', underline='single', size=10)
    base_font = Font(name='Calibri', size=10)
    wrap_align = Alignment(vertical='top', wrap_text=True)
    center_align = Alignment(horizontal='center', vertical='top')

    for r_idx, row in enumerate(rows, 2):
        fill = fill_even if r_idx % 2 == 0 else fill_odd
        for c_idx, value in enumerate(row, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=value or '')
            cell.fill = fill
            cell.border = cell_border
            # Hyperlink for website column
            if c_idx == 4 and value:
                cell.hyperlink = value
                cell.font = link_font
                cell.alignment = center_align
            elif c_idx in (2, 3, 6, 8):
                cell.font = base_font
                cell.alignment = center_align
            else:
                cell.font = base_font
                cell.alignment = wrap_align
        ws.row_dimensions[r_idx].height = 40

    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(headers))}1'

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return Response(
        out.getvalue(),
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': 'attachment; filename=events_export.xlsx'}
    )


@app.route('/api/generate-topics', methods=['POST'])
def generate_topics():
    import re
    try:
        import anthropic
    except ImportError:
        return jsonify({'error': 'anthropic package not installed'}), 500

    try:
        data = request.get_json(force=True) or {}
        topic = (data.get('topic') or '').strip()
        if not topic:
            return jsonify({'error': 'topic required'}), 400

        api_key = os.environ.get('ANTHROPIC_API_KEY', '')
        if not api_key:
            return jsonify({'error': 'ANTHROPIC_API_KEY not set — adaugă cheia în .env'}), 500

        client = anthropic.Anthropic(api_key=api_key)
        prompt = (
            f'Ești un expert în evenimente și activități pentru familii și copii din Cluj-Napoca, România.\n\n'
            f'Utilizatorul dorește să caute evenimente/workshop-uri legate de: "{topic}"\n\n'
            'Generează un JSON cu topicuri relevante și queries de căutare pentru DuckDuckGo.\n'
            'Fiecare topic să aibă:\n'
            '- "id": string unic snake_case\n'
            '- "label": etichetă scurtă (max 3 cuvinte, română)\n'
            '- "queries": listă de 4-6 queries specifice Cluj-Napoca (română + engleză)\n\n'
            'Returnează DOAR JSON valid, fără text în afara lui:\n'
            '{"topics": [{"id": "...", "label": "...", "queries": ["..."]}]}\n\n'
            'Generează 4-7 topicuri relevante.'
        )

        resp = client.messages.create(
            model='claude-haiku-4-5-20251001',
            max_tokens=1024,
            messages=[{'role': 'user', 'content': prompt}]
        )
        raw = resp.content[0].text.strip()
        raw = re.sub(r'^```(?:json)?\s*', '', raw)
        raw = re.sub(r'\s*```$', '', raw.strip())
        return jsonify(json.loads(raw))

    except Exception as e:
        return jsonify({'error': str(e)}), 500


@app.route('/api/export/xlsx/current', methods=['POST'])
def export_xlsx_current():
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    events = request.get_json(force=True) or []

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Evenimente"

    headers = ['Eveniment', 'Oraș', 'Dată', 'Website', 'Email', 'Telefon', 'Descriere', 'Salvat la']
    col_widths = [45, 16, 14, 40, 32, 18, 60, 18]

    header_font = Font(name='Calibri', bold=True, color='FFFFFF', size=11)
    header_fill = PatternFill(fill_type='solid', fgColor='1F3864')
    header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    thin = Side(style='thin', color='CCCCCC')
    cell_border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col, (h, w) in enumerate(zip(headers, col_widths), 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = cell_border
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.row_dimensions[1].height = 22

    fill_even = PatternFill(fill_type='solid', fgColor='EEF2F7')
    fill_odd  = PatternFill(fill_type='solid', fgColor='FFFFFF')
    link_font = Font(name='Calibri', color='1155CC', underline='single', size=10)
    base_font = Font(name='Calibri', size=10)
    wrap_align = Alignment(vertical='top', wrap_text=True)
    center_align = Alignment(horizontal='center', vertical='top')

    keys = ['event_name', 'city', 'date', 'webpage', 'email', 'phone', 'summary', 'timestamp']
    for r_idx, ev in enumerate(events, 2):
        row = [ev.get(k, '') or '' for k in keys]
        fill = fill_even if r_idx % 2 == 0 else fill_odd
        for c_idx, value in enumerate(row, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=value)
            cell.fill = fill
            cell.border = cell_border
            if c_idx == 4 and value:
                cell.hyperlink = value
                cell.font = link_font
                cell.alignment = center_align
            elif c_idx in (2, 3, 6, 8):
                cell.font = base_font
                cell.alignment = center_align
            else:
                cell.font = base_font
                cell.alignment = wrap_align
        ws.row_dimensions[r_idx].height = 40

    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(headers))}1'

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return Response(
        out.getvalue(),
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': 'attachment; filename=events_cautare.xlsx'}
    )


@app.route('/api/events/save', methods=['POST'])
def save_current_events():
    events = request.get_json(force=True) or []
    if not events:
        return jsonify({'ok': True, 'saved': 0})
    con = init_db()
    from email_scraper import save_events
    save_events(con, events)
    con.close()
    return jsonify({'ok': True, 'saved': len(events)})


if __name__ == '__main__':
    app.run(debug=True, threaded=True, port=5000)
